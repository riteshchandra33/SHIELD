import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/home/user/SHIELD/resume/Ritesh_Somashekar_Python_UI_Developer.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.2)
sec.top_margin = Inches(0.25)
sec.bottom_margin = Inches(0.3)
TEXT_W = Inches(8.5 - 0.4)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.0


def runs(p, text, size=11, italic=False):
    """Add text where **x** is bold."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            r = p.add_run(part)
            r.bold = i % 2 == 1
            r.italic = italic
            r.font.size = Pt(size)
    return p


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "808080")):
        b.set(qn("w:" + k), v)
    bdr.append(b)
    pPr.append(bdr)


def centered(text, size, bold=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    return p


def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(14)
    bottom_border(p)


def bullet(text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.33)
    pf.first_line_indent = Inches(-0.23)
    pf.tab_stops.add_tab_stop(Inches(0.33))
    pf.space_after = Pt(1)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("•\t").font.size = Pt(11)
    runs(p, text)


def skill(label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runs(p, f"**{label}:** {text}")


def job(company, role, right):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_W - Inches(0.05), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(company + " ")
    r.bold, r.font.size = True, Pt(12)
    r = p.add_run("| ")
    r.bold = True
    r = p.add_run(role)
    r.bold = r.italic = True
    r = p.add_run("\t" + right)
    r.bold = True


# ---------------- Header ----------------
centered("Ritesh Somashekar", 18)
centered("Python/UI Developer", 14)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Fairfax, VA | (571) 245-7770 | ").font.size = Pt(12)
r = p.add_run("ritesh.chandra33@gmail.com")
r.font.size, r.font.underline = Pt(12), True
r.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
p.add_run(" | LinkedIn").font.size = Pt(12)

# ---------------- Summary ----------------
heading("PROFESSIONAL SUMMARY")
bullet("**Python/UI Developer** with **6+ years** of experience enhancing and modernizing data platforms, building "
       "**web dashboards and APIs**, **distributed data-processing pipelines**, and **containerized applications** using "
       "**Python, Java, SQL, FastAPI, and Apache Spark**. Proven experience in **access control and authorization (OAuth 2.0)**, "
       "**job scheduling**, and **workflow orchestration (Apache Airflow, 20+ production pipelines)**.")
bullet("Strong background in **profiling and optimizing long-running compute workloads** — cutting iteration cycles from "
       "**3 days to under 18 hours**, preprocessing overhead by **70%**, and release cycles from **2 weeks to 5 days**. Experienced "
       "with **Docker, Kubernetes, and CI/CD** on **AWS** and **Azure**, delivering **HIPAA**- and **PCI-DSS**-compliant systems.")
bullet("Demonstrated success designing **user-facing dashboards and interfaces** (FastAPI, Tableau, Power BI) with clear "
       "**information hierarchy, usability, and consistent layout**. Adept at **technical design discussions, code reviews, and "
       "technical documentation**, collaborating effectively across engineering, product, research, and client teams.")

# ---------------- Skills ----------------
heading("TECHNICAL SKILLS")
skill("Programming & Scripting", "Python (pandas, NumPy, scikit-learn, PyTorch), Java, SQL (PostgreSQL, Oracle), R, Scala")
skill("Web UI, APIs & Security", "FastAPI, RESTful Microservices, API Versioning, OAuth 2.0, Dashboard Design, "
      "Tableau, Power BI, Data Visualization")
skill("Job Scheduling & Orchestration", "Apache Airflow (DAGs), HPC Job Scheduling (ORC Hopper), Kafka, AWS SQS, "
      "Event-Driven Workflows, CI/CD")
skill("Containerization & DevOps", "Docker, Kubernetes, Jenkins, Git, DVC, MLflow")
skill("Distributed Processing & Performance", "Apache Spark, PySpark, AWS EMR, Hadoop, Hive, Sqoop, ETL Pipelines, "
      "Delta Lake, Performance Tuning, Optuna")
skill("Cloud & Databases", "AWS (S3, EC2, Lambda, Glue, RDS, Athena, SageMaker), Azure (Stream Analytics, Data Services), "
      "MongoDB, Snowflake, Databricks")
skill("ML & NLP", "PyTorch, Core ML, TensorFlow, XGBoost, LightGBM, HuggingFace Transformers, BERT/BioBERT, Llama 3, SHAP")

# ---------------- Education ----------------
heading("EDUCATION")
for line in ["**Master’s in Data Analytics and Engineering** | George Mason University, Fairfax, VA, USA.",
             "**Bachelor’s in Information Science and Engineering** | Visvesvaraya Technological University, Bengaluru, Karnataka, India."]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_after = Pt(3)
    runs(p, line)

# ---------------- Experience ----------------
heading("PROFESSIONAL EXPERIENCE")
job("George Mason University", "UI/UX Developer", "Fairfax, VA | August 2024 – Present")
for b in [
    "Designed and implemented a **multi-user research dashboard** backed by a **FastAPI** and **Docker** service, structuring "
    "**information hierarchy and layout** across analysis views and serving results to users with **sub-200ms latency**.",
    "Built **Tableau** visualizations and reports for disaster-response and social-cohesion research, improving **usability and "
    "consistency** of how findings were presented to research stakeholders.",
    "Scheduled and ran **long-running compute jobs** (graph analytics and ensemble forecasting) on the **ORC Hopper HPC cluster** "
    "using **Python** and **R**, achieving an **RMSE of 0.89** on large-scale real-world datasets.",
    "Analyzed and optimized **end-to-end processing pipelines** in **Python, PyTorch, and MLflow**, automating tuning with "
    "**Optuna** to cut iteration cycles from **3 days to under 18 hours**.",
    "Built a transformer-based **document-processing pipeline** (**BioBERT**, HuggingFace) over **30K documents**, improving "
    "entity-recognition F1-score by **18%** and standardizing **15GB** across 6 datasets with **30%** less preparation time.",
    "Engineered **distributed feature pipelines** on **Apache Spark** and **AWS EMR** for large-scale telemetry data with "
    "**privacy-preserving data compliance**, and served on-device models at **sub-150ms latency** for millions of daily requests.",
    "Streamlined deployment workflows with **MLflow** and **FastAPI**, reducing release cycles from **2 weeks to 5 days**, enabling "
    "**A/B testing**, and shrinking model size by **40%** (LoRA, QAT) within **3%** of baseline accuracy.",
    "Built a **LightGBM** anomaly-detection service with threshold-based alerting under **PCI-DSS** compliance, reducing false "
    "positives by **25%**, with **SHAP**-based explainability for governance reviews.",
    "Collaborated in **design discussions and code reviews** (Git, DVC) and authored **technical documentation** and "
    "reproducibility standards, reducing onboarding time for new contributors to **under 2 days**.",
]:
    bullet(b)

job("Carelon Global Solutions", "Python Developer (Client: Accenture)", "India | January 2019 – August 2023")
for b in [
    "Developed **RESTful microservices** with **OAuth 2.0** authorization and **API versioning** for secure enterprise system "
    "communication, supporting **HIPAA-compliant** data exchange and real-time **EHR integration**.",
    "Designed and optimized **20+ data pipelines** for provider and consumer analytics, orchestrated with **Apache Airflow DAGs** "
    "and **CI/CD** on **AWS**, scheduling ingestion, feature engineering, and retraining jobs and reducing manual effort by **25%**.",
    "Streamlined **end-to-end ingestion workflows** by building **Python** automation for **NPPES** provider data, improving "
    "data accuracy by **74%** and reducing processing overhead in claims-balancing workflows.",
    "**Profiled and optimized** large-scale processing by migrating from **Hive to Apache Spark/PySpark** and re-engineering "
    "**SQL** feature pipelines, boosting model-training efficiency by **60%** and cutting preprocessing overhead by **70%**.",
    "Built scalable **distributed ETL pipelines** with **PySpark, AWS (S3, Glue, SageMaker)**, and **Delta Lake**, and web crawlers "
    "that structured **400,000+** XML/JSON documents, improving data reliability by **30%**.",
    "Containerized and deployed Python services using **Docker, Kubernetes, and Jenkins CI/CD**, improving system uptime by **20%**.",
    "Developed **Tableau** and **Power BI** dashboards and a real-time **Azure Stream Analytics** pipeline surfacing **performance "
    "metrics** and anomalies for IoT sensor data, reducing unplanned downtime by **20%**.",
    "Delivered **BERT/BioBERT** clinical entity extraction (**80% accuracy**, **55%** faster processing) and rule-based invoice "
    "classification that reduced manual error rates by **70%**; designed **MongoDB** models for high-throughput sensor data.",
    "Collaborated with cross-functional teams and **Accenture** client stakeholders on the **Seven Plus Locations** project through "
    "**technical design discussions and code reviews**, reducing project turnaround time by **15%**.",
]:
    bullet(b)

doc.core_properties.author = "Ritesh Somashekar"
doc.save(OUT)
print("saved", OUT)

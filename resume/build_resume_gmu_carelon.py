import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/home/user/SHIELD/resume/Ritesh_Somashekar_AI_ML_Engineer_GMU_Carelon.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.2)
sec.top_margin = Inches(0.25)
sec.bottom_margin = Inches(0.3)
TEXT_W = Inches(8.5 - 0.4)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.0


def runs(p, text, size=11, italic=False):
    """Add text where **x** is bold and *x* is italic."""
    for part in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not part:
            continue
        bold = part.startswith("**")
        ital = italic or (part.startswith("*") and not bold)
        r = p.add_run(part.strip("*") if (bold or ital and part.startswith("*")) else part)
        r.bold, r.italic = bold, ital
        r.font.size = Pt(size)
    return p


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "808080")):
        b.set(qn("w:" + k), v)
    bdr.append(b)
    pPr.append(bdr)


def centered(text, size, bold=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    return p


def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(14)
    bottom_border(p)


def bullet(text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.33)
    pf.first_line_indent = Inches(-0.23)
    pf.tab_stops.add_tab_stop(Inches(0.33))
    pf.space_after = Pt(1)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("•\t").font.size = Pt(11)
    runs(p, text)


def skill(label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runs(p, f"**{label}:** {text}")


def job(company, role, right):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_W - Inches(0.05), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(company + " ")
    r.bold, r.font.size = True, Pt(12)
    r = p.add_run("| ")
    r.bold = True
    r = p.add_run(role)
    r.bold = r.italic = True
    r = p.add_run("\t" + right)
    r.bold = True


# ---------------- Header ----------------
centered("Ritesh Somashekar", 18)
centered("AI & ML Engineer", 14)
LINKEDIN = "https://www.linkedin.com/in/ritesh-somashekar-0497aa182/"


def hyperlink(p, url, text, size):
    """Add a clickable link whose visible text stays readable to ATS parsers."""
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:color", "0563C1"), ("w:u", "single"), ("w:sz", str(size * 2))):
        e = OxmlElement(tag)
        e.set(qn("w:val"), val)
        rpr.append(e)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)


p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Fairfax, VA | (571) 245-7770 | ").font.size = Pt(11)
hyperlink(p, "mailto:ritesh.chandra33@gmail.com", "ritesh.chandra33@gmail.com", 11)
p.add_run(" | ").font.size = Pt(11)
hyperlink(p, LINKEDIN, "LinkedIn", 11)

# ---------------- Summary ----------------
heading("PROFESSIONAL SUMMARY")
bullet("**AI & Machine Learning Engineer** with **5+ years** of experience progressing from **data engineering** to **machine "
       "learning** and **Generative AI**, building **production RAG systems**, **LLM applications**, and **enterprise data pipelines** "
       "across healthcare and public safety. Published researcher (**ACM WSDM '26**, **IEEE PST 2026**) with hands-on expertise in "
       "**Llama 3, LlamaIndex, ChromaDB, pgvector, FAISS, Hugging Face, and PyTorch**.")
bullet("Strong background in **data engineering** with **Apache Airflow, Spark/PySpark, Snowflake, and Databricks**, and **cloud "
       "pipelines on AWS (S3, Lambda, Athena, Glue, SageMaker)** and **Azure**, with a focus on **HIPAA-compliant** healthcare data. "
       "Experienced in **MLOps** using **MLflow, Docker, Kubernetes, FastAPI,** and **CI/CD pipelines**.")
bullet("Demonstrated impact, including **RAG recall@10 raised from 61% to 88%**, a **0.81 protocol F1 with a 3.0% hallucination "
       "rate**, **99.9% pipeline reliability**, and **60% faster model training**, delivering **end-to-end AI solutions** from "
       "research to production.")

# ---------------- Skills ----------------
heading("TECHNICAL SKILLS")
skill("Programming & Scripting", "Python (NumPy, pandas, scikit-learn, PyTorch, TensorFlow), SQL, R, Scala, Java, Bash")
skill("RAG & Generative AI", "Retrieval-Augmented Generation (RAG), Chunking, Embeddings (Sentence Transformers, all-mpnet-base-v2), "
      "Metadata Filtering, Top-K Retrieval, Context Assembly, Grounding, LlamaIndex, Vector Databases (ChromaDB, pgvector, FAISS), "
      "Llama 3/3.1, Ollama, Prompt Engineering, Agentic AI, LoRA/PEFT, Quantization-Aware Training (QAT)")
skill("LLM Evaluation & Safety", "RAGAS, Recall@K, Faithfulness, Hallucination Detection, Guardrails, Ontology-Based Validation, "
      "NLI-Based Validation, Semantic Similarity Scoring, Ablation Studies, Protocol F1, Confidence Scoring")
skill("Machine Learning & NLP", "Supervised & Unsupervised Learning, Transformers, BERT, BioBERT, Hugging Face, XGBoost, LightGBM, "
      "Graph Analytics, Ensemble Models, Time Series Forecasting, Anomaly Detection, Feature Engineering, Model Explainability (SHAP)")
skill("MLOps & ML Infrastructure", "MLflow, Docker, Kubernetes, Jenkins, CI/CD, FastAPI, Git, DVC, Model Deployment")
skill("Data Engineering & Big Data", "Apache Airflow, Apache Spark, PySpark, Kafka, Databricks, Snowflake, Delta Lake, Hadoop, Hive, "
      "Sqoop, Informatica PowerCenter, Precisely Data360, SQL Server, MySQL, MongoDB")
skill("Cloud & Visualization", "AWS (S3, Lambda, Athena, Glue, SageMaker, SQS), Azure (Data Services, Stream Analytics), Tableau, "
      "Power BI (DAX), Excel (VBA)")

# ---------------- Education ----------------
heading("EDUCATION")
for line in ["**Master's in Data Analytics and Engineering** | George Mason University, Fairfax, VA, USA | GPA: 3.93/4.0",
             "**Bachelor's in Information Science and Engineering** | Visvesvaraya Technological University, Bengaluru, Karnataka, India."]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_after = Pt(3)
    runs(p, line)

# ---------------- Experience ----------------
heading("PROFESSIONAL EXPERIENCE")
job("George Mason University", "AI & ML Research Associate", "Fairfax, VA | August 2024 - Present")
for b in [
    "Co-developed **SHIELD** (published at **ACM WSDM '26**), an online conversational AI training platform for 9-1-1 call-takers "
    "built with **Fairfax County DPSC**, using **LLM agents** to generate call scenarios from CAD data, deliver **real-time "
    "metacognitive nudges**, and produce post-call reflection analytics.",
    "Co-developed **MATE-911** (published at **IEEE PST 2026**), a multimodal **RAG** system that grounds **Llama 3.1 8B** generation "
    "in validated emergency-response knowledge, producing protocol-compliant voice and text scenarios on commodity hardware with "
    "**0.81 protocol F1**, a **3.0% hallucination rate**, and a **97.6% pass rate**, with no commercial API dependency.",
    "**Led the RAG knowledge-base and data pipeline for both platforms** in **Python**, covering OCR parsing, abbreviation expansion, "
    "completeness checks, segmentation, and standardization of **160 gold-standard scenarios**; the curated knowledge base cut "
    "hallucinations from **13.0% to 3.0%**.",
    "Built the **retrieval layer** with **all-mpnet-base-v2** (768-dim) embeddings in a **ChromaDB HNSW cosine index**, attaching "
    "metadata (category, severity, tags, confidence) for **filtered top-K retrieval** and a **70% similarity relevance gate**.",
    "Engineered **prompt and context assembly** to ground **Llama 3.1 8B (Ollama)** responses in retrieved scenarios, tuning "
    "temperature, token budgets, and context windows to reach **960 ms (text)** and **1.3 s (voice)** mean latency, with fallback "
    "generation for resilience.",
    "**Led the validation layer for both platforms**, building **RAG guardrails** with a **911 ontology validator** (NENA standards, "
    "DPSC procedures), rule-based and **LLM-based semantic checks**, **NLI-based hallucination detection**, PASS/REVIEW/FAIL "
    "confidence scoring, and an **85% cosine-similarity** quality gate.",
    "Ran **RAG ablation studies** showing retrieval drives quality (**F1 0.81 vs. 0.59** and **97.6% vs. 49.7%** pass rate without "
    "it), while ontology validation acts as a safety guardrail that cuts hallucinations from **9.5% to 3.0%**.",
    "Developed a **RAG research assistant** with **LlamaIndex**, **pgvector**, and **FAISS**, optimizing chunking, embeddings, "
    "retrieval, and context assembly to raise **recall@10 from 61% to 88%**, evaluated with **RAGAS** (faithfulness, answer "
    "relevance, context precision) and served to **40+ researchers** at about **190 ms p95** latency.",
    "Applied **graph-based analytics** and **ensemble models** to large-scale disaster datasets on the **ORC Hopper HPC cluster**, "
    "forecasting community coping trends with an **RMSE of 0.89**.",
    "Implemented **model explainability** using **SHAP** for **intent classification** and **behavioral anomaly detection** models, "
    "guiding threshold tuning and model governance; the **LightGBM** anomaly detector reduced false positives by **25%**.",
    "Fine-tuned **Llama 3** with **LoRA/PEFT** and **QAT**, cutting model size by **40%** within **3%** of full-precision accuracy.",
    "Took models and RAG services from **proof of concept to production** with **MLflow** experiment tracking and model registry, "
    "**FastAPI**, **Docker**, and **CI/CD**, reducing release cycles from **2 weeks to 5 days** and enabling **A/B testing** of "
    "model and prompt versions.",
]:
    bullet(b)

job("Carelon Global Solutions (Elevance Health)", "Machine Learning Engineer",
    "Bengaluru, India | May 2022 - August 2023")
for b in [
    "Built NLP models using **BERT** and **BioBERT** for clinical entity extraction from unstructured medical text, achieving "
    "**80% accuracy** and reducing processing time by **55%**.",
    "Improved **Scikit-learn** and **XGBoost** models through feature engineering and hyperparameter tuning, increasing accuracy by "
    "**8%** and cutting training time by **15%**.",
    "Containerized ML inference services with **Docker, Kubernetes,** and **Jenkins CI/CD**, improving system uptime by **20%** "
    "and enabling real-time EHR integration via **REST APIs**.",
    "Engineered ML-ready ingestion of **400K+ XML/JSON healthcare records** into **Snowflake** and **Databricks**, accelerating "
    "model training by **60%** and reducing preprocessing time by **70%**.",
    "Built a daily **NPPES** ingestion workflow that improved ML data accuracy by **74%** and cut processing overhead by **60%**, and "
    "built **PySpark** ETL on **AWS (Glue, SageMaker)** with **Delta Lake**, improving data reliability by **30%** under **HIPAA**.",
    "Designed **20+ enterprise ETL pipelines** orchestrated with **Apache Airflow** and **AWS (S3, Lambda, Athena)**, boosting "
    "throughput by **70%** and achieving **99.9% pipeline reliability** with end-to-end lineage tracking.",
]:
    bullet(b)

job("Carelon Global Solutions (Elevance Health)", "Data Engineer", "Bengaluru, India | October 2020 - April 2022")
for b in [
    "Aligned ETL pipelines with business KPIs on the **Seven Plus Locations** project, improving claims data accuracy.",
    "Redesigned **claims-balancing ETL workflows** and automated **medical invoice processing** pipelines, reducing operational "
    "overhead by **60%** and manual errors by **70%**.",
    "Migrated large-scale processing workloads from **Hive** to **Apache Spark/PySpark**, and moved **Sqoop/Hadoop** ingestion to "
    "**Azure Data Services**, improving performance, scalability, and pipeline efficiency.",
    "Developed **RESTful microservices** with **OAuth 2.0** and API versioning, and designed **MongoDB** data models with "
    "**Python/Java** validation scripts (UDFs) to improve data quality and consistency.",
]:
    bullet(b)

# ---------------- Projects ----------------
heading("PROJECTS")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
runs(p, "**Agentic AI-Based Ticket Automation System** | SecureGPT, LLaMA, AWS SQS, JIRA, NIST-CSF")
for b in [
    "Built an **agentic AI** pipeline with **SecureGPT** and **LLaMA** on **AWS SQS** event queues that monitors cybersecurity "
    "threats and auto-generates **JIRA** tickets under **NIST-CSF**, reaching **93% automation accuracy** and cutting triage time by **70%**.",
    "Used **LLM prompt engineering** to classify and rank events, improving scalability by **50%** and cutting operational risk by **40%**.",
]:
    bullet(b)

# ---------------- Publications ----------------
heading("PUBLICATIONS")
for pub in [
    "R. S. Nayaka, **R. Somashekar**, K. B. Laskey, L. Wells II, and H. Purohit. \"SHIELD: A Framework for Online Conversational AI "
    "System for Self-Regulated Learning at Emergency Communications Centers.\" In *Proceedings of the 19th ACM International "
    "Conference on Web Search and Data Mining (WSDM '26)*, Boise, ID, 2026. doi:10.1145/3773966.3779411",
    "R. S. Nayaka, **R. Somashekar**, and H. Purohit. \"MATE-911: Multimodal AI-powered Training Environment for Next-Generation "
    "911 Dispatcher Preparedness.\" In *IEEE PST Conference*, 2026.",
]:
    bullet(pub)

doc.core_properties.author = "Ritesh Somashekar"
doc.core_properties.title = "Ritesh Somashekar - AI & ML Engineer Resume"
doc.save(OUT)
print("saved", OUT)

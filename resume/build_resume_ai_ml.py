import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/home/user/SHIELD/resume/Ritesh_Somashekar_AI_ML_Engineer.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.2)
sec.top_margin = Inches(0.25)
sec.bottom_margin = Inches(0.3)
TEXT_W = Inches(8.5 - 0.4)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.0


def runs(p, text, size=11, italic=False):
    """Add text where **x** is bold."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            r = p.add_run(part)
            r.bold = i % 2 == 1
            r.italic = italic
            r.font.size = Pt(size)
    return p


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "808080")):
        b.set(qn("w:" + k), v)
    bdr.append(b)
    pPr.append(bdr)


def centered(text, size, bold=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    return p


def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(14)
    bottom_border(p)


def bullet(text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.33)
    pf.first_line_indent = Inches(-0.23)
    pf.tab_stops.add_tab_stop(Inches(0.33))
    pf.space_after = Pt(1)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("•\t").font.size = Pt(11)
    runs(p, text)


def skill(label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runs(p, f"**{label}:** {text}")


def job(company, role, right):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_W - Inches(0.05), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(company + " ")
    r.bold, r.font.size = True, Pt(12)
    r = p.add_run("| ")
    r.bold = True
    r = p.add_run(role)
    r.bold = r.italic = True
    r = p.add_run("\t" + right)
    r.bold = True


# ---------------- Header ----------------
centered("Ritesh Somashekar", 18)
centered("AI & ML Engineer", 14)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Fairfax, VA | (571) 245-7770 | ").font.size = Pt(12)
r = p.add_run("ritesh.chandra33@gmail.com")
r.font.size, r.font.underline = Pt(12), True
r.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
p.add_run(" | LinkedIn").font.size = Pt(12)

# ---------------- Summary ----------------
heading("PROFESSIONAL SUMMARY")
bullet("**AI & Machine Learning Engineer** with **6+ years** of experience designing, deploying, and optimizing **scalable ML "
       "systems**, **NLP and LLM solutions**, and **production data pipelines** across healthcare and research domains. Proven "
       "expertise in **PyTorch, TensorFlow, and HuggingFace**, with hands-on experience in **LLM fine-tuning (LoRA, QAT)**, "
       "**agentic AI workflows**, and **low-latency model serving**.")
bullet("Strong background in **data engineering** with **Apache Spark, PySpark, Airflow, and Kafka**, and **cloud-based ML "
       "pipelines on AWS** and **Azure**, with a focus on **HIPAA-compliant** healthcare data. Experienced in **MLOps** using "
       "**MLflow, Docker, Kubernetes, FastAPI,** and **CI/CD pipelines**, enabling faster model iteration and reliable releases.")
bullet("Demonstrated impact, including **74% higher claims-data accuracy**, **60% faster model training**, and **70% fewer manual "
       "errors**, through **data-driven solutions, A/B testing,** and **model explainability (SHAP)**. Adept at partnering with "
       "cross-functional teams and client stakeholders to deliver **end-to-end ML solutions** from research to production.")

# ---------------- Skills ----------------
heading("TECHNICAL SKILLS")
skill("Programming & Scripting", "Python (NumPy, pandas, scikit-learn, TensorFlow, PyTorch), SQL (PostgreSQL, Oracle), R, Scala, Java")
skill("Machine Learning & Deep Learning", "Supervised & Unsupervised Learning, Deep Learning, Ensemble Models (XGBoost, LightGBM), "
      "Graph Analytics, Time Series Forecasting, Anomaly Detection, Feature Engineering, Hyperparameter Tuning (Optuna)")
skill("Natural Language Processing & LLMs", "NLP, Transformers, BERT, BioBERT, Llama 3, GPT-4, SecureGPT, HuggingFace, "
      "Generative AI, Agentic AI, LoRA, Quantization-Aware Training, RAG, Prompt Engineering, Vector Embeddings")
skill("MLOps & ML Infrastructure", "MLflow, Kubeflow, Airflow, Docker, Kubernetes, Jenkins, CI/CD, FastAPI, DVC, Model Deployment, "
      "Model Explainability (SHAP)")
skill("Data Engineering & Big Data", "Apache Spark, PySpark, Kafka, Hadoop, Hive, Sqoop, ETL Pipelines, Delta Lake, Snowflake, "
      "MongoDB, Web Crawling (Selenium)")
skill("Cloud Platforms", "AWS (S3, EC2, SageMaker, Lambda, Glue, Athena, SQS), Azure (Machine Learning, Stream Analytics, "
      "Data Services, Databricks)")
skill("Data Analysis & Visualization", "Tableau, Power BI (DAX), Excel (VBA), Data Visualization, Statistical Analysis, A/B Testing")

# ---------------- Education ----------------
heading("EDUCATION")
for line in ["**Master’s in Data Analytics and Engineering** | George Mason University, Fairfax, VA, USA | GPA: 3.93/4.0",
             "**Bachelor’s in Information Science and Engineering** | Visvesvaraya Technological University, Bengaluru, Karnataka, India."]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_after = Pt(3)
    runs(p, line)

# ---------------- Experience ----------------
heading("PROFESSIONAL EXPERIENCE")
job("George Mason University", "Graduate Research Assistant (Applied AI/ML)", "Fairfax, VA | August 2024 – Present")
for b in [
    "Designed **Python** and **R** ML algorithms to study **social cohesion and community coping behavior during disasters** on "
    "large-scale real-world datasets, applying **graph-based analytics** and **ensemble models** on the **ORC Hopper HPC cluster** "
    "to forecast trends with an **RMSE of 0.89**.",
    "Developed a transformer-based NLP information-extraction pipeline using **BioBERT** and **HuggingFace Transformers**, "
    "processing **30K documents** and improving entity-recognition **F1-score by 18%** over baseline spaCy models.",
    "Fine-tuned **Llama 3 LLMs** using **LoRA** and **Quantization-Aware Training (QAT)**, reducing model size by **40%** while "
    "keeping accuracy within **3%** of full-precision benchmarks.",
    "Built an **agentic AI** ticket-automation system with **SecureGPT** and **LLaMA** on **AWS SQS** to monitor security events and "
    "generate JIRA tickets under NIST-CSF, reaching **93% automation accuracy** and cutting manual triage time by **70%**.",
    "Built an **anomaly detection** system using **LightGBM** with threshold-based alerting, reducing false positives by **25%**.",
    "Built end-to-end ML pipelines with **PyTorch, MLflow,** and **Optuna**, cutting model iteration cycles from **3 days to under "
    "18 hours**, and standardized **15GB** of data across 6 datasets, reducing data-preparation time by **30%**.",
    "Deployed model inference APIs with **FastAPI** and **Docker** at **sub-200ms latency**, and streamlined MLflow-based releases "
    "from **2 weeks to 5 days** with **A/B testing** across model versions.",
    "Applied **SHAP** for model explainability, visualized findings in **R** and **Tableau**, and built reproducible workflows with "
    "**Git** and **DVC**, reducing onboarding time for new researchers to **under 2 days**.",
]:
    bullet(b)

job("Carelon Global Solutions", "Machine Learning Engineer (Client: Accenture)", "Bengaluru, India | January 2019 – August 2023")
for b in [
    "Designed and optimized **20+ ML-ready data pipelines** for provider and consumer analytics, focusing on **claims-balancing "
    "workflows** and scalable model-input generation using **Python, Apache Airflow,** and **AWS** with CI/CD.",
    "Developed Python ingestion workflows for **NPPES** provider data, improving claims and ML data accuracy by **74%** and "
    "reducing processing overhead.",
    "Built web crawlers that structured **400,000+ healthcare XML/JSON documents** for supervised learning, and engineered "
    "**SQL** and **DAX** feature pipelines that improved model training efficiency by **60%** and cut preprocessing overhead by **70%**.",
    "Designed NLP models using **BERT** and **BioBERT** for clinical entity extraction from unstructured medical text, achieving "
    "**80% accuracy** and reducing processing time by **55%** for a U.S. healthcare client.",
    "Automated **medical invoice classification** with rule-based scripting, reducing manual error rates by **70%**, and improved "
    "models with **Scikit-learn, XGBoost,** and **TensorFlow**, increasing accuracy by **8%** and cutting training time by **15%**.",
    "Built scalable **ETL pipelines** with **PySpark** and **AWS (S3, Glue, SageMaker)** on **Delta Lake**, improving data "
    "reliability by **30%** under **HIPAA**, and migrated legacy **Hive/Sqoop/Hadoop** jobs to **Apache Spark** and **Azure Data Services**.",
    "Orchestrated ingestion, feature-engineering, and model-retraining workflows with **Airflow DAGs**, reducing manual effort by **25%**.",
    "Built real-time analytics with **Azure Stream Analytics** and ML-powered **Tableau/Power BI** dashboards for anomaly "
    "detection and model-drift monitoring, reducing unplanned downtime by **20%**.",
    "Containerized ML inference services using **Docker, Kubernetes,** and **Jenkins CI/CD**, improving system uptime by **20%** "
    "and enabling real-time EHR integration via **REST APIs** secured with **OAuth 2.0**.",
    "Collaborated with Data Engineering, Product, and **Accenture** stakeholders on the **Seven Plus Locations** project to define "
    "ML use cases, KPIs, and roadmaps, reducing project turnaround time by **15%**.",
]:
    bullet(b)

doc.core_properties.author = "Ritesh Somashekar"
doc.save(OUT)
print("saved", OUT)

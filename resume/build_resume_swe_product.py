import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/home/user/SHIELD/resume/Ritesh_Somashekar_Software_Engineer_AI_Product.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.2)
sec.top_margin = Inches(0.25)
sec.bottom_margin = Inches(0.3)
TEXT_W = Inches(8.5 - 0.4)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.0


def runs(p, text, size=11, italic=False):
    """Add text where **x** is bold and *x* is italic."""
    for part in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not part:
            continue
        bold = part.startswith("**")
        ital = italic or (part.startswith("*") and not bold)
        r = p.add_run(part.strip("*") if (bold or ital and part.startswith("*")) else part)
        r.bold, r.italic = bold, ital
        r.font.size = Pt(size)
    return p


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "808080")):
        b.set(qn("w:" + k), v)
    bdr.append(b)
    pPr.append(bdr)


def centered(text, size, bold=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    return p


def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(14)
    bottom_border(p)


def bullet(text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.33)
    pf.first_line_indent = Inches(-0.23)
    pf.tab_stops.add_tab_stop(Inches(0.33))
    pf.space_after = Pt(1)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("•\t").font.size = Pt(11)
    runs(p, text)


def skill(label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runs(p, f"**{label}:** {text}")


def job(company, role, right):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_W - Inches(0.05), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(company + " ")
    r.bold, r.font.size = True, Pt(12)
    r = p.add_run("| ")
    r.bold = True
    r = p.add_run(role)
    r.bold = r.italic = True
    r = p.add_run("\t" + right)
    r.bold = True


# ---------------- Header ----------------
centered("Ritesh Somashekar", 18)
centered("Software Engineer | AI Products", 14)
LINKEDIN = "https://www.linkedin.com/in/ritesh-somashekar-0497aa182/"


def hyperlink(p, url, text, size):
    """Add a clickable link whose visible text stays readable to ATS parsers."""
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:color", "0563C1"), ("w:u", "single"), ("w:sz", str(size * 2))):
        e = OxmlElement(tag)
        e.set(qn("w:val"), val)
        rpr.append(e)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)


p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Fairfax, VA | (571) 245-7770 | ").font.size = Pt(11)
hyperlink(p, "mailto:ritesh.chandra33@gmail.com", "ritesh.chandra33@gmail.com", 11)
p.add_run(" | ").font.size = Pt(11)
hyperlink(p, LINKEDIN, "LinkedIn", 11)

# ---------------- Summary ----------------
heading("PROFESSIONAL SUMMARY")
bullet("**Software Engineer** with **5+ years** shipping **data, ML, and LLM-powered products** end to end, from research prototype "
       "to production, across healthcare, public safety, and research; published at **ACM WSDM '26** and **IEEE PST 2026**.")
bullet("Builds **AI applications** with **LLM APIs (GPT-4, Claude, Gemini, Azure OpenAI)**, **prompt engineering**, **structured "
       "outputs**, and **RAG**, and tests them for **hallucinations, safety risks, and failure modes** before users see them.")
bullet("Partners with researchers, public agencies, and product teams to turn open questions into tools people use, including "
       "**civic-tech** AI training software built with **Fairfax County 9-1-1**.")

# ---------------- Skills ----------------
heading("TECHNICAL SKILLS")
skill("Languages", "Python, SQL, Java, R")
skill("Web & Backend", "FastAPI, REST APIs, Streamlit, SQLite, PostgreSQL (pgvector), Docker, Kubernetes, CI/CD, Git")
skill("AI & LLMs", "LLM APIs (GPT-4, Claude, Gemini, Azure OpenAI), Prompt Engineering, Structured Outputs, RAG, LangChain, "
      "LlamaIndex, Vector Databases (ChromaDB, Pinecone, FAISS), Llama 3, Ollama, Agentic AI")
skill("AI Evaluation & Safety", "LLM Evaluation, RAGAS, Hallucination Detection, Guardrails, Ablation Studies, Failure-Mode "
      "Analysis, A/B Testing")
skill("Data & Research Tools", "Pandas, Apache Airflow, Spark/PySpark, Snowflake, Databricks, MLflow, Tableau, Power BI")
skill("Machine Learning & Cloud", "PyTorch, Scikit-learn, Hugging Face, BERT, XGBoost, AWS (S3, Lambda, Glue, SageMaker, SQS), Azure")

# ---------------- Education ----------------
heading("EDUCATION")
for line in ["**Master's in Data Analytics and Engineering** | George Mason University, Fairfax, VA, USA | GPA: 3.93/4.0",
             "**Bachelor's in Information Science and Engineering** | Visvesvaraya Technological University, Bengaluru, India"]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_after = Pt(3)
    runs(p, line)

# ---------------- Experience ----------------
# Each bullet follows STAR: situation/task -> action (with tools) -> measurable result.
heading("PROFESSIONAL EXPERIENCE")
job("George Mason University", "AI & ML Research Associate", "Fairfax, VA | August 2024 - Present")
for b in [
    "To scale 9-1-1 call-taker training beyond instructor role-play, co-developed **SHIELD** with **Fairfax County DPSC**, a web app "
    "where **LLM agents** simulate calls and coach trainees live; published at **ACM WSDM '26**.",
    "To move SHIELD from prototype to pilot, built its **Python/Streamlit** interface with **SQLite** accounts, course enrollment, "
    "and session tracking, iterating on feedback from instructors and agency testing sessions.",
    "To generate safe training calls without commercial APIs, co-developed **MATE-911**, a **RAG** application on **Llama 3.1 8B** "
    "reaching **0.81 protocol F1** and a **3.0% hallucination rate**; published at **IEEE PST 2026**.",
    "To give instructors trustworthy content, led the **data pipeline** turning unstructured call records into **160 gold-standard "
    "scenarios**, cutting hallucinations from **13.0% to 3.0%**.",
    "To stop unsafe outputs before they reached trainees, led the **validation layer** with **protocol guardrails** and "
    "**hallucination checks**, lowering hallucinations from **9.5% to 3.0%**.",
    "To learn what actually drove quality, ran **ablation studies** that exposed failure modes, showing **semantic search** "
    "retrieval lifted the pass rate from **49.7% to 97.6%**.",
    "To keep live sessions responsive, engineered **prompt and context assembly** for **Llama 3.1 8B**, delivering responses in "
    "under **1 second** for text and **1.3 seconds** for voice.",
    "To speed up literature review, built a **RAG research assistant** with **LlamaIndex**, **pgvector**, and a **FastAPI** "
    "backend, raising **recall@10 from 61% to 88%** for **40+ researchers** at **190 ms p95** latency.",
    "To compare LLM providers, integrated **GPT-4, Claude,** and **Gemini** APIs through **LangChain** with **structured outputs**, "
    "raising answer accuracy by **16%** and cutting hallucinations by **19%**.",
    "To ship research prototypes reliably, set up **MLflow**, **Docker**, and **CI/CD**, cutting release cycles from **2 weeks to "
    "5 days** with **A/B testing** of model and prompt versions.",
    "To support disaster-response research, built **Python** and **R** analysis workflows on the **ORC Hopper HPC cluster**, "
    "forecasting community coping trends with an **RMSE of 0.89**.",
]:
    bullet(b)

job("Carelon Global Solutions (Elevance Health)", "Machine Learning Engineer", "Bengaluru, India | May 2022 - August 2023")
for b in [
    "To speed up member support, built a **Generative AI** assistant with **Azure OpenAI**, **LangChain**, and **Pinecone** for "
    "**15K+ daily queries**, cutting resolution time by **38%** and raising ticket deflection by **25%**.",
    "To automate clinical text review, built **BERT/BioBERT** extraction models on unstructured medical notes, reaching **80% "
    "accuracy** and cutting processing time by **55%**.",
    "To serve models inside clinical applications, deployed services with **Docker, Kubernetes,** and **Jenkins CI/CD** behind "
    "**REST APIs**, improving uptime by **20%** and enabling real-time EHR integration.",
    "To keep production models accurate over time, set up **monitoring**, **drift detection**, and automated retraining, "
    "sustaining **over 95% accuracy** for **12 months** across **200+ end users**.",
    "To speed up model development, built ingestion of **400K+ XML/JSON healthcare records** into **Snowflake** and "
    "**Databricks**, accelerating training by **60%** and cutting preprocessing time by **70%**.",
    "To feed products reliable data, designed **20+ pipelines** with **Airflow, PySpark,** and **AWS (S3, Lambda, Glue)**, "
    "raising throughput by **70%** and reaching **99.9% reliability** under **HIPAA**.",
]:
    bullet(b)

job("Carelon Global Solutions (Elevance Health)", "Data Engineer", "Bengaluru, India | October 2020 - April 2022")
for b in [
    "To reduce manual claims reconciliation, redesigned **claims-balancing workflows** in **Python** and **SQL**, cutting "
    "operational overhead by **60%**.",
    "To eliminate costly invoice errors, automated **medical invoice processing** with rule-based classification, reducing "
    "manual errors by **70%**.",
    "To give business teams faster answers, delivered self-service **Tableau** and **Power BI** reporting tools, cutting "
    "reporting latency by **40%**.",
]:
    bullet(b)

# ---------------- Projects ----------------
heading("PROJECTS")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
runs(p, "**Agentic AI Ticket Automation System** | SecureGPT, LLaMA, AWS SQS, JIRA")
bullet("To cut security-ticket triage time, built an **agentic AI** system with **SecureGPT** and **LLaMA** on **AWS SQS** that "
       "classifies threats and auto-creates **JIRA** tickets, reaching **93% accuracy** and **70%** faster triage.")

# ---------------- Publications ----------------
heading("PUBLICATIONS")
for pub in [
    "R. S. Nayaka, **R. Somashekar**, K. B. Laskey, L. Wells II, H. Purohit. \"SHIELD: A Framework for Online Conversational AI "
    "System for Self-Regulated Learning at Emergency Communications Centers.\" *ACM WSDM '26*. doi:10.1145/3773966.3779411",
    "R. S. Nayaka, **R. Somashekar**, H. Purohit. \"MATE-911: Multimodal AI-powered Training Environment for Next-Generation "
    "911 Dispatcher Preparedness.\" *IEEE PST 2026*.",
]:
    bullet(pub)

doc.core_properties.author = "Ritesh Somashekar"
doc.core_properties.title = "Ritesh Somashekar - Software Engineer, AI Products Resume"
doc.core_properties.keywords = "Software Engineer, Product, LLM, AI Applications, RAG, Full-Stack, Research Tools"
doc.save(OUT)
print("saved", OUT)

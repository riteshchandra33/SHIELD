"""Builds the cover letter for the Collective Intelligence Project Software Engineer, Product role."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/home/user/SHIELD/resume/Ritesh_Somashekar_Cover_Letter_CIP.docx"
LINKEDIN = "https://www.linkedin.com/in/ritesh-somashekar-0497aa182/"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.9)
sec.top_margin = sec.bottom_margin = Inches(0.7)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(11)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.space_after = Pt(8)
st.paragraph_format.line_spacing = 1.08


def hyperlink(p, url, text):
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:color", "0563C1"), ("w:u", "single")):
        e = OxmlElement(tag)
        e.set(qn("w:val"), val)
        rpr.append(e)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)


def centered(text, size, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(size)


def rule(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "4"), ("color", "808080")):
        b.set(qn("w:" + k), v)
    bdr.append(b)
    pPr.insert(0, bdr)


def para(text, after=8, justify=True):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    return p


# Header (matches the resume)
centered("Ritesh Somashekar", 18)
centered("Software Engineer | AI Products", 13)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
p.add_run("Fairfax, VA | (571) 245-7770 | ")
hyperlink(p, "mailto:ritesh.chandra33@gmail.com", "ritesh.chandra33@gmail.com")
p.add_run(" | ")
hyperlink(p, LINKEDIN, "LinkedIn")
rule(p)

para("October 23, 2026", after=10, justify=False)
para("Hiring Team\nCollective Intelligence Project\nRe: Software Engineer, Product", after=10, justify=False)
para("Dear Collective Intelligence Project Hiring Team,", justify=False)

para("I am applying for the Software Engineer, Product role. For the past two years at George Mason University, I have been "
     "turning research questions about AI into software that real people use. CIP's mission, bringing public judgment into how "
     "AI is built and evaluated, is the work I want to do next.")

para("With Fairfax County's 9-1-1 center, I co-developed SHIELD, a web app where LLM agents simulate emergency calls and coach "
     "trainees in real time (published at ACM WSDM '26). I built its Python and Streamlit interface, with accounts, course "
     "enrollment, and session tracking, and improved it through instructor feedback and agency testing sessions. On a related "
     "system, MATE-911 (IEEE PST 2026), I led the data pipeline and the validation layer: guardrails and hallucination checks "
     "that brought the hallucination rate down to 3.0%, plus ablation studies that showed which components failed and why.")

para("That evaluation work is why Weval stands out to me. Whether a model's answer is good enough depends on who is asking and "
     "what is at stake; for a 9-1-1 trainee, a fluent but wrong answer is a safety risk. I would welcome the chance to build "
     "tools that let communities and domain experts set those standards themselves, and to help Global Dialogues turn what "
     "people say into findings that labs and policymakers can act on.")

para("Before graduate school, I spent nearly three years at Carelon Global Solutions, moving from data engineering to machine "
     "learning. I shipped an Azure OpenAI and LangChain assistant serving 15,000+ daily member queries and ran production "
     "services with monitoring and CI/CD. At GMU, I also built a RAG research assistant used by 40+ researchers and integrated "
     "the GPT-4, Claude, and Gemini APIs with structured outputs.")

para("My interface work so far has been in Python rather than TypeScript and React, and I am ready to close that gap quickly; "
     "moving across stacks is how I grew from data pipelines to LLM products. I would value the chance to discuss how I could "
     "help CIP turn research into products people can rely on.")

para("Sincerely,", after=2, justify=False)
para("Ritesh Somashekar", after=0, justify=False)

doc.core_properties.author = "Ritesh Somashekar"
doc.core_properties.title = "Ritesh Somashekar - Cover Letter - Software Engineer, Product"
doc.save(OUT)
print("saved", OUT)

"""Builds the AI/ML Systems Analyst/Engineer resume (Resume 2 layout, tuned for AI/Gen AI testing roles)."""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/home/user/SHIELD/resume/Ritesh_Somashekar_AI_ML_Systems_Engineer.docx"
LINK_BLUE = RGBColor(0x05, 0x63, 0xC1)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.28)
sec.top_margin = Inches(0.4)
sec.bottom_margin = Inches(0.35)
TEXT_W = Inches(8.5 - 0.56)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.0


def runs(p, text, size=11):
    """Add text where **x** is bold."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            r = p.add_run(part)
            r.bold = i % 2 == 1
            r.font.size = Pt(size)
    return p


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", "000000")):
        b.set(qn("w:" + k), v)
    bdr.append(b)
    pPr.insert(0, bdr)


def centered(text, size):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(size)


def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(14)
    bottom_border(p)


def bullet(text, after=2):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.25)
    pf.first_line_indent = Inches(-0.25)
    pf.tab_stops.add_tab_stop(Inches(0.25))
    pf.space_after = Pt(after)
    p.add_run("•\t")
    runs(p, text)


def skill(label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    runs(p, f"**{label}:** {text}")


def job(left, right):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(f"{left}\t{right}")
    r.bold, r.font.size = True, Pt(12.5)


# ---------------- Header ----------------
centered("Ritesh Somashekar", 18)
centered("AI/ML Systems Analyst & Engineer", 14)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Fairfax, VA | ")
for i, (txt, link) in enumerate([("somashekarritesh@gmail.com", True), (" | (571) 245-7770 | ", False), ("LinkedIn", True)]):
    r = p.add_run(txt)
    if link:
        r.font.underline, r.font.color.rgb = True, LINK_BLUE

# ---------------- Summary ----------------
heading("PROFESSIONAL SUMMARY")
bullet("**AI/ML Systems Engineer** with **6+ years of experience** building, **testing, and validating** end-to-end **AI/ML, "
       "Generative AI, RAG, NLP, and data platforms**, including **3+ years of cloud AI** on **AWS (SageMaker, Bedrock)** and "
       "**Azure OpenAI**, with deep expertise in **Python, SQL (Oracle), PySpark, Kafka, LangChain, and LlamaIndex**.", 4)
bullet("Proven impact in AI quality, including improving **RAG recall from 61% to 88%**, **reducing LLM hallucination rate by 19%**, "
       "sustaining **95%+ production model accuracy** through drift detection, and **cutting reporting errors by 40%** with automated "
       "data quality frameworks.", 4)
bullet("Specialized in **test strategy for non-deterministic Gen AI systems** (accuracy, relevance, faithfulness, safety, and bias), "
       "**BDD (Cucumber/Gherkin), Selenium, Pytest**, API and system testing, and **LLM evaluation frameworks (RAGAS, DeepEval)**; "
       "experienced in **Agile** delivery with **Jira and Confluence**, communicating findings to technical and business stakeholders.", 4)

# ---------------- Skills ----------------
heading("TECHNICAL SKILLS")
skill("Programming & Query Languages", "Python, SQL (Oracle, PL/SQL, PostgreSQL, SQL Server), NoSQL, Gherkin")
skill("Testing & QA Automation", "BDD (Cucumber, Behave, pytest-bdd), Selenium WebDriver, Pytest, Postman, REST API Testing, "
      "System, Integration & Regression Testing, Manual & Exploratory Testing, Test Harness & Test Strategy Design")
skill("AI/LLM Evaluation & Testing", "RAGAS, DeepEval, LLM-as-a-Judge, Golden Datasets, Hallucination Detection, Bias & Fairness "
      "Evaluation, Model Drift Monitoring, Prompt Regression Testing, A/B Testing, Embedding & Retrieval Validation")
skill("LLM & Generative AI", "Large Language Models, Prompt Engineering, Retrieval-Augmented Generation (RAG), Vector Databases "
      "(FAISS, Pinecone, pgvector), LangChain, LlamaIndex, LoRA/PEFT Fine-Tuning, Structured Outputs, Embeddings")
skill("Machine Learning & Frameworks", "ML Model Lifecycle, Deep Learning, NLP, Anomaly Detection, Time Series Forecasting, "
      "Feature Engineering, PyTorch, TensorFlow, Scikit-learn, Hugging Face Transformers, XGBoost, LightGBM, Pandas, NumPy")
skill("MLOps & Model Deployment", "Docker, Kubernetes, CI/CD Pipelines, GitHub Actions, MLflow, Kubeflow, Model Monitoring, "
      "Experiment Tracking, Model Serving (FastAPI, Flask), Terraform")
skill("Cloud Platforms", "AWS (Bedrock, SageMaker, SageMaker Feature Store, S3, Lambda, EC2, Glue, EKS, IAM), Azure (Azure ML, "
      "Azure OpenAI, Data Factory, Data Lake), Snowflake, Databricks")
skill("Data Engineering", "Apache Spark, PySpark, Apache Airflow, Apache Kafka, Apache Flink, dbt, ETL/ELT Pipelines, Data Lakes, Hadoop")
skill("Collaboration & Tools", "Jira, Confluence, Agile/Scrum, Git, Linux/Unix, Jupyter Notebook, ELK Stack, Power BI, Tableau")

# ---------------- Experience ----------------
heading("PROFESSIONAL EXPERIENCE")
job("Apple | AI/ML Engineer", "CA | May 2025 – Present")
for b in [
    "Designed an **evaluation harness for non-deterministic LLM outputs** using golden datasets, **RAGAS**, and **LLM-as-a-judge** "
    "scoring for accuracy, relevance, faithfulness, and safety, used as a release gate for customer-facing Gen AI features.",
    "Deployed production LLM applications for customer service automation using **prompt engineering** and **structured outputs**, "
    "adding **hallucination detection** and output-validation checks that reduced hallucination rate **19%** and response latency **20%**.",
    "Built **RAG-based retrieval systems** using **LangChain**, embedding models, and **FAISS** to ground **GPT-4, Claude, and Gemini** "
    "outputs in verified sources, validating retrieval relevance and **boosting response accuracy by 16%**.",
    "Ran **systematic model evaluation and A/B testing** to benchmark LLM quality pre/post-deployment, including **prompt regression "
    "and bias evaluation**; logged defects in **Jira** and documented test strategy and findings in **Confluence**.",
    "Automated **BDD regression suites (pytest-bdd, Gherkin)** and **API tests** for LLM inference endpoints, integrated into **CI/CD** "
    "with **MLflow, Docker, and GitHub Actions** to validate every model and prompt release.",
    "Engineered **real-time data pipelines** with **Apache Spark, Flink, and Kafka** for NLP models, adding automated data validation "
    "checks and improving pipeline **throughput by 26%**.",
    "Fine-tuned transformer models with **PyTorch**, applying **LoRA, PEFT, and knowledge distillation** for intent classification and "
    "ticket routing, cutting model size **36%** while holding **87%+ task accuracy** against regression benchmarks.",
]:
    bullet(b)

job("George Mason University | Graduate Research Assistant (Applied ML)", "VA | August 2024 – May 2025")
for b in [
    "Developed a **RAG research assistant** with **LlamaIndex**, prompt engineering, **pgvector**, and **FAISS**, validating "
    "embeddings, retrieval, and context assembly to improve **recall@10 from 61% to 88%**.",
    "Established an **LLM evaluation framework** using **RAGAS, MLflow, Git, and DVC** to measure faithfulness, answer relevance, and "
    "context precision, reducing iteration cycles from **3 days to 16 hours**; protocols adopted by **two follow-on LLM studies**.",
    "Deployed the RAG application with **FastAPI and Docker**, building **automated API tests (Pytest)** and **Selenium UI tests** for "
    "grounded, cited responses, supporting **40+ researchers** at approximately **190ms p95 latency**.",
    "Built a biomedical ML data pipeline using **Python, SQL, Pandas, AWS S3, and DVC**, curating **30K+ publications** with data "
    "validation, deduplication, metadata tagging, and reproducible train/eval splits.",
    "Fine-tuned **BioBERT** with **PyTorch and Hugging Face** for biomedical NER using **Optuna** and **PEFT/LoRA**, improving "
    "**F1 by 18 points** over spaCy across **6K+ annotated sentences**.",
]:
    bullet(b)

job("Accenture | Machine Learning Engineer", "India | September 2020 – July 2023")
for b in [
    "Delivered a **Generative AI** customer support solution with **Azure OpenAI LLMs** and **RAG** for a retail client, handling "
    "**15K+ daily queries**; participated in **business requirements reviews** and designed end-to-end test scenarios with product owners.",
    "Integrated **RAG workflows** with **LangChain**, vector embeddings, and **Pinecone**, validating retrieval and response quality "
    "to reduce resolution time **38%** and increase ticket deflection **25%**.",
    "Wrote **complex Oracle SQL** queries and automated query processing across **Oracle, Databricks, and SageMaker Feature Store** "
    "to validate feature data, model metadata, and prediction outputs before each release.",
    "Built **Selenium and Cucumber (BDD)** automation for web and mobile support workflows, and orchestrated **MLOps** with **Docker, "
    "Kubernetes, MLflow, and AWS SageMaker**, accelerating the model deployment lifecycle from **2 weeks to 3 days**.",
    "Engineered **PySpark and Databricks** pipelines processing **500GB+** of structured and unstructured data, reducing processing "
    "time by **45%**, and developed **XGBoost** forecasting models improving accuracy by **22%** across **10+ categories**.",
    "Established **model monitoring, drift detection**, and automated retraining to maintain **95%+ production accuracy** over "
    "**12 months** for **200+ end users**, working in Agile sprints tracked in **Jira** and **Confluence**.",
]:
    bullet(b)

job("Honeywell | Data Engineer", "India | January 2019 – September 2020")
for b in [
    "Engineered **ETL/ELT pipelines** for industrial IoT analytics using **Python, PySpark, and Apache Spark**, ingesting telemetry "
    "from **8+ source systems** via REST APIs and processing **2M+ records daily** into **Azure Data Lake**.",
    "Implemented **data quality and system testing frameworks** (null checks, duplicate detection, schema validation, referential "
    "integrity, anomaly detection), reducing downstream reporting errors by **40%** across **3 plant sites**.",
    "Automated batch workflows with **Azure Data Factory** and **Apache Airflow**, including source-to-target mapping validation and "
    "monitoring across **15+ daily jobs**, cutting issue detection time to **under 15 minutes**.",
    "Built reusable **SQL Server and Oracle** transformation and reconciliation logic with **dbt**, improving data prep efficiency by "
    "**35%**, and designed a **star-schema** warehouse that cut report turnaround by **60%** across **5 teams**.",
    "Managed production deployments with **Git, CI/CD, and Linux/Shell scripting**, archived telemetry in **AWS S3**, and partnered "
    "with Power BI developers on **3 on-time dashboard releases**.",
]:
    bullet(b)

# ---------------- Education ----------------
heading("EDUCATION")
for line in ["**Master’s in Data Analytics and Engineering** | George Mason University, Fairfax, VA, USA.",
             "**Bachelor’s in Information Science and Engineering** | Visvesvaraya Technological University, Bengaluru, Karnataka, India."]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    runs(p, line)

doc.core_properties.author = "Ritesh Somashekar"
doc.save(OUT)
print("saved", OUT)
